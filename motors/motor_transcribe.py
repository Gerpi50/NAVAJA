"""
motors/motor_transcribe.py — Transcribir un audio o un video con Whisper.

Usa el paquete `transcripcion/` (Whisper): elige entre procesador y placa
de video y transcribe en castellano con tildes. Un archivo puede durar
horas, y por eso el resultado se devuelve partido en tramos con su minuto y
segundo: así se puede escuchar y corregir tramo por tramo.

Todo funciona sin internet, con los modelos que ya están en el equipo.
"""
import hashlib
import json
import os
import subprocess
import time
import types
from typing import Callable, List, Optional


# Los modelos que se ofrecen, en este orden. El resto de los que conoce el
# motor no se muestra: tiny y base no alcanzan para una declaración, y el
# grande (large-v3) pesa 2,9 GB, pide 16 GB de memoria y casi no mejora al
# Turbo, así que no viaja con el programa.
MODELOS = [
    ("medium",         "Medio — recomendado: muy buena precisión"),
    ("large-v3-turbo", "Turbo — muy buena precisión, conviene en grabaciones largas"),
    ("small",          "Chico — rápido, para una primera pasada"),
]
MODELO_POR_DEFECTO = "medium"

DISPOSITIVOS = [
    ("auto", "Automático (recomendado)"),
    ("cpu",  "Procesador"),
    ("cuda", "Placa de video"),
]

# Le marca a Whisper el estilo: castellano rioplatense, con tildes y signos.
# Va CON tildes a propósito: Whisper imita la ortografía de este texto.
INDICACION = (
    "Transcripción literal en español de una entrevista grabada. Se usa "
    "voseo rioplatense, con tildes, signos de interrogación y puntuación "
    "normal. Por ejemplo: ¿Cómo te llamás? Sí, ayer a la tarde estaba en la "
    "casa de mi abuela."
)

SIN_VOZ = "(no se escucha voz en este tramo)"


class Cancelado(Exception):
    """El usuario cortó la transcripción."""


# ── Disponibilidad ───────────────────────────────────────────────────────

def motivo_no_disponible() -> str:
    """"" si se puede transcribir; si no, la explicación para el usuario."""
    import importlib.util
    try:
        falta = importlib.util.find_spec("whisper") is None
    except Exception:
        falta = True
    if falta:
        return ("En esta instalación falta el motor de transcripción "
                "(Whisper).\nAbrí el programa con «Abrir_App (codigo "
                "fuente).bat», que usa el Python que sí lo tiene.")
    from motors.motor_video_compress import FFMPEG
    if not FFMPEG:
        return ("Falta FFmpeg: colocá ffmpeg.exe y ffprobe.exe en la carpeta "
                "bin/ del programa.")
    return ""


def modelos_disponibles() -> list:
    """[(nombre, etiqueta, está_en_el_equipo)] de los modelos que se ofrecen."""
    from transcripcion import transcriber as wa
    return [(nombre, etiqueta, os.path.isfile(wa.model_path(nombre)))
            for nombre, etiqueta in MODELOS]


# Memoria que conviene que tenga la PC (en total) para cada modelo. Son
# valores prudentes: con menos puede andar, pero Windows empieza a usar el
# disco como memoria y una transcripción de minutos pasa a tardar horas.
MEMORIA_PARA = {"small": 4, "medium": 8, "large-v3-turbo": 8, "large-v3": 16}


def memoria_gb() -> float:
    """Memoria RAM instalada en la PC, en GB (0 si no se pudo averiguar)."""
    try:
        import ctypes

        class _Estado(ctypes.Structure):
            _fields_ = [("largo", ctypes.c_ulong), ("carga", ctypes.c_ulong),
                        ("total", ctypes.c_ulonglong),
                        ("libre", ctypes.c_ulonglong),
                        ("resto", ctypes.c_ulonglong * 5)]

        estado = _Estado()
        estado.largo = ctypes.sizeof(_Estado)
        if ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(estado)):
            return estado.total / (1024 ** 3)
    except Exception:
        pass
    return 0.0


def _alcanza(modelo: str, memoria: float) -> bool:
    # Una PC "de 8 GB" informa 7,7 o 7,9: por eso el margen.
    return memoria <= 0 or memoria >= MEMORIA_PARA.get(modelo, 8) - 0.6


def modelo_sugerido() -> str:
    """El modelo que conviene en esta PC: Medio, o Chico si hay poca memoria."""
    return MODELO_POR_DEFECTO if _alcanza(MODELO_POR_DEFECTO, memoria_gb()) \
        else "small"


def aviso_memoria(modelo: str) -> str:
    """"" si la memoria alcanza para ese modelo; si no, el aviso."""
    memoria = memoria_gb()
    if _alcanza(modelo, memoria):
        return ""
    return (f"Esta PC tiene {memoria:.0f} GB de memoria y este modelo anda "
            f"bien con {MEMORIA_PARA.get(modelo, 8)} GB o más: puede tardar "
            f"muchísimo o cortarse. Acá conviene el modelo «Chico».")


def describir_dispositivo(preferencia: str = "auto") -> str:
    """Con qué se va a transcribir y por qué (consulta la placa: tarda)."""
    from transcripcion import transcriber as wa
    _disp, motivo = wa.pick_device(preferencia)
    return f"{wa.describe_device(preferencia)}: {motivo}."


# ── Audio para escuchar ──────────────────────────────────────────────────

def extraer_audio(src: str, dst_wav: str) -> str:
    """Saca el audio a un WAV mono de 32 kHz.

    Sirve para dos cosas: es lo que lee Whisper y es lo que suena al revisar
    (un WAV se puede recorrer y cambiar de velocidad sin saltos, cosa que no
    pasa con todos los formatos de origen).
    """
    from transcripcion import entorno
    from motors.motor_video_compress import FFMPEG
    entorno.preparar_entorno()
    os.makedirs(os.path.dirname(dst_wav) or ".", exist_ok=True)
    r = subprocess.run(
        [FFMPEG, "-y", "-v", "error", "-i", src, "-vn", "-sn",
         "-map", "0:a:0", "-ac", "1", "-ar", "32000", "-c:a", "pcm_s16le",
         dst_wav],
        capture_output=True, timeout=3600,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if r.returncode != 0 or not os.path.isfile(dst_wav) \
            or os.path.getsize(dst_wav) < 1000:
        raise RuntimeError(
            "No pude leer el audio de ese archivo.\n"
            + r.stderr.decode(errors="replace")[-300:])
    return dst_wav


# ── Transcribir ──────────────────────────────────────────────────────────

def transcribir(
    wav: str,
    modelo: str = MODELO_POR_DEFECTO,
    dispositivo: str = "auto",
    progreso_cb: Optional[Callable[[float], None]] = None,
    log_cb: Optional[Callable[[str], None]] = None,
    cancel_cb: Optional[Callable[[], bool]] = None,
) -> dict:
    """Transcribe y devuelve los tramos.

    {"tramos": [{"inicio": s, "fin": s, "texto": str}, …],
     "modelo": str, "dispositivo": "cpu"|"cuda", "segundos": float}

    progreso_cb recibe de 0 a 1. Si cancel_cb() da True, corta con Cancelado.
    """
    sesion = Sesion(modelo, dispositivo, log_cb)
    try:
        if cancel_cb and cancel_cb():
            raise Cancelado()
        return sesion.transcribir(wav, progreso_cb, cancel_cb)
    finally:
        sesion.cerrar()


class Sesion:
    """El modelo cargado una vez, para transcribir uno o varios archivos.

    Cargar el modelo lleva varios segundos: con una carpeta de notas de voz
    cortas, hacerlo por cada archivo sería más lento que transcribirlas.
    """

    def __init__(self, modelo: str = MODELO_POR_DEFECTO,
                 dispositivo: str = "auto",
                 log_cb: Optional[Callable[[str], None]] = None):
        from transcripcion import entorno
        from transcripcion import transcriber as wa
        entorno.preparar_entorno()      # Whisper llama a «ffmpeg» por el PATH
        self._wa = wa
        self.modelo = modelo
        self._log = log_cb or (lambda _texto: None)
        self._motor = wa.Transcriber(modelo, log_cb=self._log,
                                     device_pref=dispositivo)
        self._motor.ensure_model()      # avisa en castellano si falta el .pt

    @property
    def dispositivo(self) -> str:
        return self._motor.device

    def cerrar(self) -> None:
        self._motor.release()

    def transcribir(
        self,
        wav: str,
        progreso_cb: Optional[Callable[[float], None]] = None,
        cancel_cb: Optional[Callable[[], bool]] = None,
        tramos_cb: Optional[Callable[[List[dict]], None]] = None,
    ) -> dict:
        """Transcribe un WAV; devuelve lo mismo que `transcribir()`.

        tramos_cb recibe los tramos a medida que salen (cada medio minuto de
        grabación, más o menos), para ir mostrándolos. Son provisorios: los
        definitivos son los que se devuelven al final.
        """
        wa = self._wa
        t0 = time.monotonic()

        # El módulo, no la función: «whisper.transcribe» a secas es la función
        # del mismo nombre, que tapa al módulo.
        import importlib
        wt = importlib.import_module("whisper.transcribe")

        # Whisper no avisa por dónde va: lo único que hace es mover una barra
        # de consola (tqdm). Se la cambia por una propia, que avisa el avance
        # y de paso es el lugar donde se puede cortar.
        class _Barra:
            def __init__(self, total=None, **_kw):
                self.total = total or 1
                self.hecho = 0

            def __enter__(self):
                return self

            def __exit__(self, *_exc):
                return False

            def update(self, n=1):
                self.hecho += n
                if progreso_cb:
                    progreso_cb(max(0.0, min(1.0, self.hecho / self.total)))
                if cancel_cb and cancel_cb():
                    raise Cancelado()

        # Tampoco avisa qué va entendiendo. Pero por cada medio minuto llama a
        # la función que le pone el momento a cada palabra: ahí se lo pesca.
        tiempos_original = wt.add_word_timestamps

        def tiempos_y_aviso(*args, **kwargs):
            tiempos_original(*args, **kwargs)
            if tramos_cb:
                segmentos = kwargs.get("segments") or (args[0] if args else [])
                try:
                    nuevos = _armar_tramos(segmentos, wa._tidy)
                    if nuevos:
                        tramos_cb(nuevos)
                except Exception:
                    pass                # mostrar el avance nunca corta el trabajo

        opciones = dict(
            language="es", task="transcribe",
            fp16=(self._motor.device == "cuda"),
            initial_prompt=INDICACION,
            # Sin esto, en grabaciones largas Whisper se engancha repitiendo
            # la última frase durante minutos.
            condition_on_previous_text=False,
            no_speech_threshold=0.6,
            # Con el momento de cada palabra se pueden armar tramos de una
            # oración. Sin esto, el modelo Turbo devuelve bloques de 30
            # segundos, que no sirven para revisar escuchando.
            word_timestamps=True,
            verbose=False,
        )

        tqdm_original = wt.tqdm
        wt.tqdm = types.SimpleNamespace(tqdm=_Barra)
        wt.add_word_timestamps = tiempos_y_aviso
        try:
            try:
                resultado = self._motor._model.transcribe(wav, **opciones)
            except Cancelado:
                raise
            except Exception as exc:
                if self._motor.device == "cuda" and wa._is_oom(exc):
                    self._log("La placa de video se quedó sin memoria: sigo "
                              "con el procesador.")
                    self._motor.release()
                    self._motor = wa.Transcriber(
                        self.modelo, log_cb=self._log, device_pref="cpu")
                    self._motor.ensure_model()
                    opciones["fp16"] = False
                    resultado = self._motor._model.transcribe(wav, **opciones)
                else:
                    raise
        finally:
            wt.tqdm = tqdm_original
            wt.add_word_timestamps = tiempos_original

        juntos = _armar_tramos(resultado.get("segments") or [], wa._tidy)
        return {"tramos": juntos, "modelo": self.modelo,
                "dispositivo": self._motor.device,
                "segundos": time.monotonic() - t0}


# Un tramo se cierra al terminar una oración; si la oración es eterna, o la
# persona hace una pausa larga, también.
_TRAMO_MAX_S = 14.0
_PAUSA_S = 1.2
_FIN_DE_ORACION = (".", "?", "!", "…")


def _armar_tramos(segmentos: list, prolijar) -> List[dict]:
    """Convierte lo que devuelve Whisper en tramos de más o menos una oración."""
    tramos: List[dict] = []

    def cerrar(palabras):
        texto = prolijar("".join(p["word"] for p in palabras).strip())
        if texto:
            tramos.append({"inicio": float(palabras[0]["start"]),
                           "fin": float(palabras[-1]["end"]),
                           "texto": texto})

    for seg in segmentos:
        # Sobre silencio Whisper a veces inventa una frase: si está casi
        # seguro de que no hay voz, se deja constancia en vez de inventar.
        sin_voz = seg.get("no_speech_prob", 0) > 0.85
        if sin_voz or not (seg.get("text") or "").strip():
            tramos.append({"inicio": float(seg.get("start", 0.0)),
                           "fin": float(seg.get("end", 0.0)),
                           "texto": SIN_VOZ})
            continue
        palabras = seg.get("words") or []
        if not palabras:                # sin tiempos por palabra: tal cual
            tramos.append({"inicio": float(seg.get("start", 0.0)),
                           "fin": float(seg.get("end", 0.0)),
                           "texto": prolijar(seg["text"].strip())})
            continue
        actual = []
        for p in palabras:
            if actual and p["start"] - actual[-1]["end"] > _PAUSA_S:
                cerrar(actual)
                actual = []
            actual.append(p)
            largo = p["end"] - actual[0]["start"]
            fin_oracion = p["word"].rstrip().endswith(_FIN_DE_ORACION)
            if fin_oracion or largo > _TRAMO_MAX_S:
                cerrar(actual)
                actual = []
        if actual:
            cerrar(actual)

    # Los tramos sin voz seguidos se juntan en uno solo.
    juntos: List[dict] = []
    for t in tramos:
        if juntos and t["texto"] == SIN_VOZ and juntos[-1]["texto"] == SIN_VOZ:
            juntos[-1]["fin"] = t["fin"]
        else:
            juntos.append(t)
    return juntos


# ── Guardar el resultado ─────────────────────────────────────────────────

def marca(segundos: float) -> str:
    s = int(max(0, segundos))
    if s >= 3600:
        return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}"
    return f"{s // 60:02d}:{s % 60:02d}"


def _marca_srt(segundos: float) -> str:
    ms = int(round(max(0.0, segundos) * 1000))
    return (f"{ms // 3600000:02d}:{ms % 3600000 // 60000:02d}:"
            f"{ms % 60000 // 1000:02d},{ms % 1000:03d}")


def quien(tramo: dict, nombres: Optional[list]) -> str:
    """El nombre de quien habla en ese tramo ("" si no se sabe)."""
    persona = tramo.get("persona")
    if not nombres or persona is None or not (0 <= persona < len(nombres)):
        return ""
    return nombres[persona]


def nombres_utiles(tramos: list, nombres: Optional[list]) -> Optional[list]:
    """Los nombres, sólo si en la grabación habla más de uno.

    Con una sola voz no tiene sentido encabezar cada renglón con
    «Interlocutor 1:».
    """
    if not nombres:
        return None
    hablan = {t.get("persona") for t in tramos
              if t.get("persona") is not None
              and 0 <= t["persona"] < len(nombres)}
    return nombres if len(hablan) > 1 else None


def _intervenciones(tramos: list, nombres: Optional[list]) -> list:
    """Junta los tramos seguidos de una misma persona: [(nombre, texto)]."""
    partes: list = []
    for t in tramos:
        if t["texto"] == SIN_VOZ:
            continue
        nombre = quien(t, nombres)
        if partes and partes[-1][0] == nombre:
            partes[-1][1] += " " + t["texto"]
        else:
            partes.append([nombre, t["texto"]])
    return [(nombre, texto) for nombre, texto in partes]


def texto_completo(tramos: list, con_tiempos: bool = True,
                   nombres: Optional[list] = None) -> str:
    """La transcripción entera como texto.

    Con `nombres` (los participantes) cada renglón dice quién habla. Sin
    marcas de tiempo sale como un diálogo: un párrafo por intervención.
    """
    nombres = nombres_utiles(tramos, nombres)
    if con_tiempos:
        renglones = []
        for t in tramos:
            nombre = quien(t, nombres)
            renglones.append(f"[{marca(t['inicio'])}] "
                             + (f"{nombre}: " if nombre else "") + t["texto"])
        return "\n".join(renglones)
    partes = _intervenciones(tramos, nombres)
    if any(nombre for nombre, _texto in partes):
        return "\n\n".join((f"{nombre}: " if nombre else "") + texto
                           for nombre, texto in partes)
    # Sin participantes ni tiempos: texto corrido.
    return " ".join(texto for _nombre, texto in partes)


def guardar(tramos: list, ruta: str, con_tiempos: bool = True,
            titulo: str = "", nombres: Optional[list] = None) -> str:
    """Guarda en .txt, .srt o .docx según la extensión de `ruta`."""
    ext = os.path.splitext(ruta)[1].lower()
    os.makedirs(os.path.dirname(ruta) or ".", exist_ok=True)
    nombres = nombres_utiles(tramos, nombres)
    if ext == ".srt":
        partes = []
        for i, t in enumerate(tramos, start=1):
            nombre = quien(t, nombres)
            partes.append(f"{i}\n{_marca_srt(t['inicio'])} --> "
                          f"{_marca_srt(t['fin'])}\n"
                          + (f"{nombre}: " if nombre else "")
                          + f"{t['texto']}\n")
        with open(ruta, "w", encoding="utf-8") as f:
            f.write("\n".join(partes))
    elif ext == ".docx":
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        estilo = doc.styles["Normal"]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(12)
        if titulo:
            doc.add_heading(titulo, level=1)
        if con_tiempos:
            for t in tramos:
                nombre = quien(t, nombres)
                p = doc.add_paragraph()
                p.add_run(f"[{marca(t['inicio'])}] "
                          + (f"{nombre}: " if nombre else "")).bold = True
                p.add_run(t["texto"])
        else:
            for nombre, texto in _intervenciones(tramos, nombres):
                p = doc.add_paragraph()
                if nombre:
                    p.add_run(f"{nombre}: ").bold = True
                p.add_run(texto)
        doc.save(ruta)
    else:
        with open(ruta, "w", encoding="utf-8") as f:
            if titulo:
                f.write(titulo + "\n\n")
            f.write(texto_completo(tramos, con_tiempos, nombres) + "\n")
    return ruta


# ── El informe: la transcripción con el hash de cada archivo ─────────────
#
# Abre con un párrafo que dice qué se transcribió, consigna el hash SHA-256
# de lo que se analizó y no agrega nada del programa (ni carátula, ni firma,
# ni mención de la herramienta).
#
# Cada archivo es un dict: {"nombre", "sha256", "tamano" (bytes), "duracion"
# (s), "es_video", "tramos", "nombres"}.

SIN_VOZ_EN_ARCHIVO = "(no se registra voz en la grabación)"


def _tipo_de(archivos: list) -> str:
    hay_audio = any(not a.get("es_video") for a in archivos)
    hay_video = any(a.get("es_video") for a in archivos)
    if hay_audio and hay_video:
        return "audio y video"
    return "video" if hay_video else "audio"


def encabezado_informe(archivos: list, carpeta: str = "") -> str:
    """El párrafo de apertura, con el hash si se trata de un solo archivo."""
    if len(archivos) == 1:
        a = archivos[0]
        return (
            f"Transcripción: se analizó el archivo de {_tipo_de(archivos)} "
            f"“{a['nombre']}”, de {marca(a.get('duracion', 0))} de duración, "
            f"cuyo hash SHA-256 es {a.get('sha256', '')}. Su contenido es "
            "el siguiente:")
    donde = f" contenidos en la carpeta “{carpeta}”" if carpeta else ""
    return (
        f"Transcripción: se analizaron {len(archivos)} archivos de "
        f"{_tipo_de(archivos)}{donde}, cuyos hash SHA-256 se consignan en "
        "cada caso. Su contenido es el siguiente:")


def _renglones(archivo: dict, con_tiempos: bool) -> list:
    """La transcripción de un archivo como [(parte en negrita, texto)]."""
    tramos = archivo.get("tramos") or []
    nombres = nombres_utiles(tramos, archivo.get("nombres"))
    if con_tiempos:
        salida = []
        for t in tramos:
            nombre = quien(t, nombres)
            salida.append((f"[{marca(t['inicio'])}] "
                           + (f"{nombre}: " if nombre else ""), t["texto"]))
        return salida or [("", SIN_VOZ_EN_ARCHIVO)]
    partes = _intervenciones(tramos, nombres)
    if not partes:
        return [("", SIN_VOZ_EN_ARCHIVO)]
    if any(nombre for nombre, _texto in partes):
        return [(f"{nombre}: " if nombre else "", texto)
                for nombre, texto in partes]
    return [("", " ".join(texto for _nombre, texto in partes))]


def _titulo_de_archivo(numero: int, archivo: dict) -> str:
    from transcripcion import human_size
    return (f"{numero}) Archivo “{archivo['nombre']}” "
            f"({human_size(archivo.get('tamano', 0))}, duración "
            f"{marca(archivo.get('duracion', 0))})")


def guardar_informe(archivos: list, ruta: str, con_tiempos: bool = True,
                    carpeta: str = "") -> str:
    """Escribe el informe en .odt, .docx o .txt según la extensión."""
    ext = os.path.splitext(ruta)[1].lower()
    os.makedirs(os.path.dirname(ruta) or ".", exist_ok=True)
    encabezado = encabezado_informe(archivos, carpeta)
    varios = len(archivos) > 1
    corte = encabezado.find(":") + 1       # «Transcripción:» va en negrita

    if ext == ".odt":
        from transcripcion.odt_writer import OdtReport
        doc = OdtReport()
        doc.add_rich_paragraph([(encabezado[:corte], "Negrita"),
                                (encabezado[corte:], None)])
        for numero, archivo in enumerate(archivos, start=1):
            doc.add_empty()
            if varios:
                doc.add_rich_paragraph(
                    [(_titulo_de_archivo(numero, archivo), "Negrita")])
                doc.add_paragraph(f"Hash SHA-256: {archivo.get('sha256', '')}")
            for negrita, texto in _renglones(archivo, con_tiempos):
                doc.add_rich_paragraph([(negrita, "Negrita"), (texto, None)])
        doc.save(ruta)
    elif ext == ".docx":
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        estilo = doc.styles["Normal"]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(12)
        p = doc.add_paragraph()
        p.add_run(encabezado[:corte]).bold = True
        p.add_run(encabezado[corte:])
        for numero, archivo in enumerate(archivos, start=1):
            if varios:
                doc.add_paragraph()
                doc.add_paragraph().add_run(
                    _titulo_de_archivo(numero, archivo)).bold = True
                doc.add_paragraph(f"Hash SHA-256: {archivo.get('sha256', '')}")
            for negrita, texto in _renglones(archivo, con_tiempos):
                p = doc.add_paragraph()
                if negrita:
                    p.add_run(negrita).bold = True
                p.add_run(texto)
        doc.save(ruta)
    else:
        lineas = [encabezado]
        for numero, archivo in enumerate(archivos, start=1):
            lineas.append("")
            if varios:
                lineas.append(_titulo_de_archivo(numero, archivo))
                lineas.append(f"Hash SHA-256: {archivo.get('sha256', '')}")
            salto = "\n" if con_tiempos else "\n\n"
            lineas.append(salto.join(
                negrita + texto
                for negrita, texto in _renglones(archivo, con_tiempos)))
        with open(ruta, "w", encoding="utf-8") as f:
            f.write("\n".join(lineas) + "\n")
    return ruta


# ── Trabajo guardado, para seguir otro día ───────────────────────────────

def carpeta_trabajos() -> str:
    base = os.getenv("LOCALAPPDATA") or os.path.expanduser("~")
    return os.path.join(base, "Navaja", "transcripciones")


def ruta_trabajo(src: str) -> str:
    """Un archivo de trabajo por archivo de origen (ruta + tamaño + fecha)."""
    try:
        st = os.stat(src)
        firma = f"{os.path.abspath(src).lower()}|{st.st_size}|{int(st.st_mtime)}"
    except OSError:
        firma = os.path.abspath(src).lower()
    clave = hashlib.sha1(firma.encode("utf-8")).hexdigest()[:20]
    return os.path.join(carpeta_trabajos(), f"{clave}.trabajo.json")


def guardar_trabajo(src: str, datos: dict) -> str:
    ruta = ruta_trabajo(src)
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    datos = dict(datos, origen=os.path.abspath(src),
                 guardado=time.strftime("%Y-%m-%d %H:%M:%S"))
    provisorio = ruta + ".tmp"
    with open(provisorio, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=1)
    os.replace(provisorio, ruta)        # nunca queda un archivo a medio escribir
    return ruta


def cargar_trabajo(src: str) -> Optional[dict]:
    try:
        with open(ruta_trabajo(src), "r", encoding="utf-8") as f:
            datos = json.load(f)
        return datos if datos.get("tramos") else None
    except (OSError, ValueError):
        return None


def borrar_trabajo(src: str) -> None:
    try:
        os.remove(ruta_trabajo(src))
    except OSError:
        pass


# ── Autoprueba (verifica que el paquete portable puede transcribir) ──────

def autoprueba(args: list) -> int:
    """Transcribe un audio sin abrir la ventana y deja el resultado escrito.

    Lo usa `main.py --probar-transcripcion <audio> <informe.txt> [modelo]`.
    Sirve para comprobar, en una PC nueva, que la carpeta trae todo lo que
    hace falta. Devuelve 0 si salió bien.
    """
    import sys
    import tempfile
    import traceback

    audio = args[0] if args else ""
    informe = args[1] if len(args) > 1 else os.path.join(
        tempfile.gettempdir(), "prueba_transcripcion.txt")
    modelo = args[2] if len(args) > 2 else MODELO_POR_DEFECTO
    lineas = [f"Empaquetado: {bool(getattr(sys, 'frozen', False))}",
              f"Ejecutable: {sys.executable}"]
    codigo = 1
    try:
        motivo = motivo_no_disponible()
        lineas.append(
            f"Motor disponible: {'sí' if not motivo else 'NO: ' + motivo}")
        from transcripcion import transcriber as wa
        lineas.append("Carpetas de modelos: " + " | ".join(wa.model_dirs()))
        for nombre, _etiqueta, esta in modelos_disponibles():
            lineas.append(f"  modelo {nombre}: "
                          f"{wa.model_path(nombre) if esta else 'NO ESTÁ'}")
        lineas.append("Dispositivo: " + describir_dispositivo("auto"))
        if audio:
            wav = os.path.join(tempfile.mkdtemp(prefix="prueba_tr_"), "a.wav")
            extraer_audio(audio, wav)
            r = transcribir(wav, modelo, "auto")
            lineas.append(f"Transcripto con «{r['modelo']}» en "
                          f"{r['segundos']:.1f} s ({r['dispositivo']}), "
                          f"{len(r['tramos'])} tramos:")
            codigo = 0 if r["tramos"] else 1
            # El reconocedor de voces es una pieza aparte: se prueba también.
            from motors import motor_voces as mv
            sin_voces = mv.motivo_no_disponible()
            nombres = None
            if sin_voces:
                lineas.append("Reconocer quién habla: NO: " + sin_voces)
                codigo = 1
            else:
                personas = mv.asignar(r["tramos"], mv.reconocer(wav, 0))
                nombres = mv.nombres_por_defecto(personas)
                lineas.append(f"Reconocer quién habla: sí, {personas} "
                              "voz/voces distintas")
            lineas.append(texto_completo(r["tramos"], True, nombres))
        else:
            codigo = 0 if not motivo else 1
        lineas.append("RESULTADO: " + ("TODO BIEN" if codigo == 0 else "FALLÓ"))
    except Exception:
        lineas.append("RESULTADO: FALLÓ")
        lineas.append(traceback.format_exc())
    try:
        with open(informe, "w", encoding="utf-8") as f:
            f.write("\n".join(lineas) + "\n")
    except OSError:
        pass
    return codigo
